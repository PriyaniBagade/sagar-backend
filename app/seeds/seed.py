"""
Seed script — loads ports from data/Ports Data.xlsx and vessel classes.
Wipes existing ports + vessel_classes rows first so re-running is safe.

Run with:
    python -m app.seeds.seed
"""

import re
from datetime import date
from pathlib import Path
import openpyxl
import csv

from app.config.database import SessionLocal, engine, Base
from app.models.port import Port, PortType
from app.models.vessel_class import VesselClass
from app.models.route_distance import RouteDistance

import app.models.port  # noqa
import app.models.vessel_class  # noqa
import app.models.route_distance  # noqa

Base.metadata.create_all(bind=engine)

DATA_FILE = Path(__file__).parent.parent.parent / "data" / "Ports Data.xlsx"
ROUTES_FILE = Path(__file__).parent.parent.parent / "data" / "route_distances.csv"
PORT_CHARGES_FILE = (
    Path(__file__).parent.parent.parent / "data" / "port_charges_mvp.csv"
)


def _load_port_charges() -> dict[str, float]:
    """Return {port_name: flat_charge_usd} from port_charges_mvp.csv.
    Normalises all dash variants to a plain hyphen so xlsx vs CSV differences don't break matching.
    """

    def _norm(s: str) -> str:
        return (
            s.replace("\u2014", "-")
            .replace("\u2013", "-")
            .replace("\u2012", "-")
            .strip()
        )

    charges: dict[str, float] = {}
    if not PORT_CHARGES_FILE.exists():
        print(
            f"WARNING: port_charges_mvp.csv not found — port_charges_flat_usd will be null."
        )
        return charges
    with open(PORT_CHARGES_FILE, newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            name = _norm(row["port_name"])
            try:
                charges[name] = float(row["flat_charge_usd"])
            except (ValueError, KeyError):
                pass
    return charges


def _parse_draft(raw) -> float | None:
    """Handle values like 14.5, '14.5/16.50', '13+ tide', 'N/A'."""
    if raw is None:
        return None
    s = str(raw).strip()
    if s.upper() in ("N/A", "N/a", ""):
        return None
    # "14.5/16.50" — take the conservative (smaller) value
    if "/" in s:
        parts = [p.strip() for p in s.split("/")]
        nums = []
        for p in parts:
            m = re.search(r"[\d.]+", p)
            if m:
                nums.append(float(m.group()))
        return min(nums) if nums else None
    # "13+ tide" — extract the number
    m = re.search(r"[\d.]+", s)
    return float(m.group()) if m else None


def _parse_float(raw) -> float | None:
    """Return float or None for N/A, N/a, text strings."""
    if raw is None:
        return None
    s = str(raw).strip()
    if s.upper() in ("N/A", "N/A", "N/a", ""):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _parse_int(raw) -> int | None:
    f = _parse_float(raw)
    return int(f) if f is not None else None


def seed():
    db = SessionLocal()
    try:
        # --- wipe existing data ---
        db.query(Port).delete()
        db.query(VesselClass).delete()
        db.query(RouteDistance).delete()
        db.commit()
        print("Cleared existing ports, vessel classes, and route distances.")

        # --- vessel classes (Baltic standard values) ---
        vessel_classes = [
            VesselClass(
                vessel_class="Capesize",
                index="BCI",
                dwt=180000,
                draft=18.2,
                loa=290,
                beam=45.0,
            ),
            VesselClass(
                vessel_class="Panamax",
                index="BPI",
                dwt=82500,
                draft=14.43,
                loa=229,
                beam=32.25,
            ),
            VesselClass(
                vessel_class="Supramax",
                index="BSI",
                dwt=63500,
                draft=13.418,
                loa=199.98,
                beam=32.24,
            ),
            VesselClass(
                vessel_class="Handysize",
                index="BHSI",
                dwt=38200,
                draft=10.538,
                loa=180,
                beam=29.8,
            ),
        ]
        db.add_all(vessel_classes)

        # --- ports from xlsx ---
        port_charges = _load_port_charges()

        def _norm_name(s: str) -> str:
            return (
                s.replace("\u2014", "-")
                .replace("\u2013", "-")
                .replace("\u2012", "-")
                .strip()
            )

        wb = openpyxl.load_workbook(DATA_FILE)
        ws = wb[wb.sheetnames[0]]  # "Ports Data " (note trailing space)

        ports = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            (
                name,
                country,
                port_type_raw,
                draft_raw,
                loa_raw,
                beam_raw,
                berths_raw,
                _,
            ) = row

            if not name or not port_type_raw:
                continue

            try:
                port_type = PortType[port_type_raw.strip().upper()]
            except KeyError:
                print(f"  Skipping unknown port type '{port_type_raw}' for {name}")
                continue

            port_name = str(name).strip()
            ports.append(
                Port(
                    name=port_name,
                    country=str(country).strip() if country else "Unknown",
                    type=port_type,
                    max_draft=_parse_draft(draft_raw),
                    max_loa=_parse_float(loa_raw),
                    max_beam=_parse_float(beam_raw),
                    berths=_parse_int(berths_raw),
                    port_charges_flat_usd=port_charges.get(_norm_name(port_name)),
                    last_verified=date.today(),
                )
            )

        db.add_all(ports)
        db.commit()

        # --- route distances from CSV ---
        routes = []
        if ROUTES_FILE.exists():
            with open(ROUTES_FILE, newline="", encoding="cp1252") as f:
                reader = csv.DictReader(f)
                for row in reader:
                    lp = row["loading_port"].strip()
                    dp = row["discharge_port"].strip()
                    if not lp or not dp:
                        continue
                    try:
                        dist = float(row["distance_nm"])
                    except (ValueError, KeyError):
                        continue
                    routes.append(
                        RouteDistance(
                            loading_port=lp,
                            discharge_port=dp,
                            distance_nm=dist,
                            route_type=row["route_type"].strip(),
                        )
                    )
            db.add_all(routes)
            db.commit()
            print(f"Seeded {len(routes)} route distances from route_distances.csv.")
        else:
            print(
                f"WARNING: route_distances.csv not found at {ROUTES_FILE} — skipping."
            )

        print(
            f"Seeded {len(vessel_classes)} vessel classes and {len(ports)} ports from Ports Data.xlsx."
        )

    except Exception as e:
        db.rollback()
        raise e
    finally:
        db.close()


if __name__ == "__main__":
    seed()
